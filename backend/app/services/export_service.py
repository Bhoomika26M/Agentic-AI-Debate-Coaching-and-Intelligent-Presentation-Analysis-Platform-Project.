import io
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

class ExportService:
    """Export Engine: Generates professional PDF and multi-tab Excel reports."""

    def generate_excel_report(self, report_data: Dict[str, Any]) -> io.BytesIO:
        wb = Workbook()
        
        # Style helpers
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        sub_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        bold_font = Font(name="Arial", size=10, bold=True)
        normal_font = Font(name="Arial", size=10)

        # Tab 1: Scores Overview
        ws_scores = wb.active
        ws_scores.title = "Scores"
        ws_scores.append(["DebateAI — Performance Evaluation Report"])
        ws_scores["A1"].font = title_font
        ws_scores.append([])
        
        scores_headers = ["Metric", "Score / 100", "Category Weight", "Status"]
        ws_scores.append(scores_headers)
        for col_num in range(1, len(scores_headers) + 1):
            cell = ws_scores.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        score_rows = [
            ["Argument Quality", report_data.get("argument_quality", 78.0), "30%", "Proficient"],
            ["Evidence Usage", report_data.get("evidence_usage", 72.0), "20%", "Satisfactory"],
            ["Logical Consistency", report_data.get("logical_consistency", 80.0), "20%", "Advanced"],
            ["Rebuttal Effectiveness", report_data.get("rebuttal_effectiveness", 75.0), "15%", "Proficient"],
            ["Communication Skills", report_data.get("communication_skills", 82.0), "15%", "Advanced"],
            ["Overall Weighted Performance", report_data.get("overall_score", 77.3), "100%", "Mastery Level 3"]
        ]
        for row in score_rows:
            ws_scores.append(row)

        for row in ws_scores.iter_rows(min_row=4, max_row=4 + len(score_rows) - 1, min_col=1, max_col=4):
            for cell in row:
                cell.font = normal_font
                if cell.column in [2, 3]:
                    cell.alignment = Alignment(horizontal="center")

        # Tab 2: Debate History
        ws_debates = wb.create_sheet(title="Debate History")
        ws_debates.append(["Round", "Speaker", "Argument Transcript", "Strength Score", "Reasoning Score"])
        for col_num in range(1, 6):
            cell = ws_debates.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        history = report_data.get("rounds_history", [
            {"round": 1, "speaker": "User", "text": "Adopting AI in education personalizes learning paths and addresses achievement gaps.", "strength": 82.0, "reasoning": 78.0},
            {"round": 1, "speaker": "AI Opponent", "text": "Over-reliance on automation risks cognitive atrophy and algorithmic bias in grading.", "strength": 85.0, "reasoning": 82.0}
        ])
        for item in history:
            ws_debates.append([item.get("round", 1), item.get("speaker", "User"), item.get("text", ""), item.get("strength", 80.0), item.get("reasoning", 75.0)])

        # Tab 3: Fallacies Detected
        ws_fallacies = wb.create_sheet(title="Fallacies")
        ws_fallacies.append(["Fallacy Type", "Confidence", "Problematic Statement", "Correction Recommendation"])
        for col_num in range(1, 5):
            cell = ws_fallacies.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        fallacies = report_data.get("fallacies", [
            {"name": "Hasty Generalization", "confidence": "85%", "statement": "Everyone always ignores the baseline cost.", "correction": "Use probabilistic qualifying phrases with demographic citations."}
        ])
        for f in fallacies:
            ws_fallacies.append([f.get("name", ""), f.get("confidence", "80%"), f.get("statement", ""), f.get("correction", "")])

        # Tab 4: Presentation Analytics
        ws_pres = wb.create_sheet(title="Presentation Analytics")
        ws_pres.append(["Metric Name", "Value", "Target Benchmark", "Assessment"])
        for col_num in range(1, 5):
            cell = ws_pres.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        ws_pres.append(["Words Per Minute (WPM)", report_data.get("wpm", 142.0), "130 - 165 WPM", "Balanced"])
        ws_pres.append(["Filler Words Count", report_data.get("filler_count", 6), "< 5 per 2 min", "Attention Needed"])
        ws_pres.append(["Confidence Score", report_data.get("confidence_score", 76.0), "> 75 / 100", "Strong"])
        ws_pres.append(["Clarity Score", report_data.get("clarity_score", 80.0), "> 75 / 100", "Superior"])

        # Tab 5: Recommendations
        ws_recs = wb.create_sheet(title="Recommendations")
        ws_recs.append(["Category", "Coaching Insight", "Suggested Drill"])
        for col_num in range(1, 4):
            cell = ws_recs.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        ws_recs.append(["Evidence Integration", "Cite peer-reviewed longitudinal studies with methodology.", "Evidence Attribution Drill"])
        ws_recs.append(["Fallacy Prevention", "Avoid universal quantifiers ('always', 'never').", "Qualifying Language Exercise"])
        ws_recs.append(["Vocal Delivery", "Introduce deliberate 1.5s pauses instead of filler words.", "2-Minute Controlled Pacing Drill"])

        # Auto-adjust column widths
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = col[0].column_letter
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def generate_pdf_report(self, report_data: Dict[str, Any]) -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=22, textColor=colors.HexColor("#1E3A8A"), spaceAfter=6)
        subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=11, textColor=colors.HexColor("#4B5563"), spaceAfter=14)
        h2_style = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=14, textColor=colors.HexColor("#1E40AF"), spaceBefore=12, spaceAfter=8)
        body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=10, textColor=colors.HexColor("#1F2937"), leading=14)
        bold_body = ParagraphStyle('BoldBody', parent=styles['Normal'], fontSize=10, fontName="Helvetica-Bold", textColor=colors.HexColor("#111827"))

        elements = []

        # Header Title
        elements.append(Paragraph("DebateAI — Official Performance Report", title_style))
        topic_title = report_data.get("topic", "AI in Modern Education & Society")
        elements.append(Paragraph(f"Topic: <b>{topic_title}</b> | Position: <b>{report_data.get('position', 'For')}</b>", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3B82F6"), spaceAfter=15))

        # Overall Weighted Score Highlight Table
        overall_score = report_data.get("overall_score", 77.3)
        score_data = [
            [Paragraph("<b>OVERALL PERFORMANCE SCORE</b>", bold_body), Paragraph(f"<b><font size=16 color='#1E40AF'>{overall_score}/100</font></b>", bold_body)]
        ]
        score_table = Table(score_data, colWidths=[350, 180])
        score_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#EFF6FF")),
            ('BOX', (0,0), (-1,-1), 1.5, colors.HexColor("#3B82F6")),
            ('ALIGN', (1,0), (1,0), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('PADDING', (0,0), (-1,-1), 10),
        ]))
        elements.append(score_table)
        elements.append(Spacer(1, 14))

        # Section: Weighted Scoring Breakdown
        elements.append(Paragraph("Weighted Performance Criteria (Project Specification)", h2_style))
        criteria_table_data = [
            ["Criterion", "Weight", "Score", "Evaluation"],
            ["Argument Quality", "30%", f"{report_data.get('argument_quality', 78.0)} / 100", "Strong thematic coherence"],
            ["Evidence Usage", "20%", f"{report_data.get('evidence_usage', 72.0)} / 100", "Good empirical reference; needs method"],
            ["Logical Consistency", "20%", f"{report_data.get('logical_consistency', 80.0)} / 100", "High deductive integrity"],
            ["Rebuttal Effectiveness", "15%", f"{report_data.get('rebuttal_effectiveness', 75.0)} / 100", "Pungent defensive response"],
            ["Communication Skills", "15%", f"{report_data.get('communication_skills', 82.0)} / 100", "Clear articulation & pacing"]
        ]
        crit_table = Table(criteria_table_data, colWidths=[160, 60, 90, 220])
        crit_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1E3A8A")),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0,0), (-1,0), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E5E7EB")),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F9FAFB")]),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        elements.append(crit_table)
        elements.append(Spacer(1, 14))

        # Section: Coaching Recommendations & Action Plan
        elements.append(Paragraph("Personalized Coaching Insights", h2_style))
        strongest = report_data.get("strongest_argument", "Clear formulation of the primary premise establishing public benefit.")
        weakest = report_data.get("weakest_argument", "Vulnerability in handling the transition cost counterargument.")
        elements.append(Paragraph(f"<b>Strongest Argument:</b> {strongest}", body_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(f"<b>Area for Improvement:</b> {weakest}", body_style))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(f"<b>Recommended Practice Drill:</b> {report_data.get('next_exercise', 'Practice identifying and countering logical fallacies')}", body_style))
        elements.append(Spacer(1, 14))

        # Build Document
        doc.build(elements)
        buffer.seek(0)
        return buffer

export_service = ExportService()
